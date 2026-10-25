import datetime
import logging
import os
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd

logger = logging.getLogger(__name__)


def parse_upcoming_event(calendar_data):
    """Parses yfinance calendar data to find the nearest upcoming corporate event."""
    today = datetime.date.today()
    event_date = "TBD"
    event_name = "Board Meeting"

    if calendar_data is None:
        return event_date, event_name

    if isinstance(calendar_data, pd.DataFrame):
        try:
            if not calendar_data.empty:
                calendar_data = calendar_data.iloc[:, 0].to_dict()
            else:
                return event_date, event_name
        except Exception:
            return event_date, event_name

    if not isinstance(calendar_data, dict):
        return event_date, event_name

    norm_cal = {str(k).lower().strip(): v for k, v in calendar_data.items()}

    earnings_val = norm_cal.get("earnings date") or norm_cal.get("earnings_date")
    if earnings_val:
        target_date = (
            earnings_val[0]
            if isinstance(earnings_val, list) and len(earnings_val) > 0
            else earnings_val
        )
        if hasattr(target_date, "date"):
            target_date = target_date.date()
        elif isinstance(target_date, str):
            try:
                target_date = datetime.datetime.strptime(
                    target_date, "%Y-%m-%d"
                ).date()
            except ValueError:
                target_date = None

        if isinstance(target_date, datetime.date) and target_date >= today:
            return target_date.strftime("%Y-%m-%d"), "Earnings Call"

    ex_div = norm_cal.get("ex-dividend date") or norm_cal.get("ex_dividend_date")
    if ex_div:
        if hasattr(ex_div, "date"):
            ex_div = ex_div.date()
        elif isinstance(ex_div, str):
            try:
                ex_div = datetime.datetime.strptime(ex_div, "%Y-%m-%d").date()
            except ValueError:
                ex_div = None

        if isinstance(ex_div, datetime.date) and ex_div >= today:
            return ex_div.strftime("%Y-%m-%d"), "Dividend Ex-Date"

    return event_date, event_name


def get_market_cap_tier_inr(market_cap):
    """Categorizes asset size based on SEBI-equivalent INR thresholds."""
    if market_cap is None or pd.isna(market_cap):
        return "N/A"
    try:
        market_cap = float(market_cap)
    except (ValueError, TypeError):
        return "N/A"

    if market_cap >= 1000000000000:
        return "Mega Cap"
    elif market_cap >= 850000000000:
        return "Large Cap"
    elif market_cap >= 220000000000:
        return "Mid Cap"
    else:
        return "Small Cap"


def format_holding_duration(delta) -> str:
    """Formats timedelta into 'X days X hours'."""
    if pd.isna(delta):
        return ""
    total_seconds = int(delta.total_seconds())
    if total_seconds < 0:
        total_seconds = 0
    days = total_seconds // 86400
    hours = (total_seconds % 86400) // 3600
    return f"{days} days {hours} hours"


def enrich_open_journal_df(
    open_journal_df: pd.DataFrame, df_scored_results: pd.DataFrame
) -> pd.DataFrame:
    """Enriches open_journal_df with scoring and grading columns from df_scored_results."""
    if open_journal_df is None or open_journal_df.empty:
        return open_journal_df
    if df_scored_results is None or df_scored_results.empty:
        return open_journal_df

    scored_map = {}
    for _, r in df_scored_results.iterrows():
        t = str(r["Ticker"]).replace(".NS", "").upper()
        scored_map[t] = {
            "Market CAP": r.get("Market_Cap_Tier", ""),
            "ATH (P-66%)": r.get("ATH_Price_Grade", ""),
            "ATP (P-75%)": r.get("ATH_Profit_Grade", ""),
            "Nifty 500 [OTP] (P-82%)": r.get("Nifty_Alpha_Grade", ""),
            "Sector Index (P-82%)": r.get("Sector_Alpha_Grade", ""),
            "REVISION SCORE": r.get("Review_Score", ""),
            "ACTION": r.get("ACTION", ""),
        }

    enriched_rows = []
    for _, row in open_journal_df.iterrows():
        inst = str(row["Instrument"]).replace(".NS", "").upper()
        d = row.to_dict()
        scores = scored_map.get(
            inst,
            {
                "Market CAP": "",
                "ATH (P-66%)": "",
                "ATP (P-75%)": "",
                "Nifty 500 [OTP] (P-82%)": "",
                "Sector Index (P-82%)": "",
                "REVISION SCORE": "",
                "ACTION": "",
            },
        )
        d.update(scores)
        enriched_rows.append(d)

    return pd.DataFrame(enriched_rows).sort_values('Instrument')


def enrich_closed_journal_df(
    closed_journal_df: pd.DataFrame, df_scored_results: pd.DataFrame
) -> pd.DataFrame:
    """Enriches closed_journal_df with scoring and grading columns from df_scored_results."""
    if closed_journal_df is None or closed_journal_df.empty:
        return closed_journal_df
    if df_scored_results is None or df_scored_results.empty:
        return closed_journal_df

    scored_map = {}
    for _, r in df_scored_results.iterrows():
        t = str(r["Ticker"]).replace(".NS", "").upper()
        scored_map[t] = {
            "Market CAP": r.get("Market_Cap_Tier", ""),
            "ATH (P-66%)": r.get("ATH_Price_Grade", ""),
            "ATP (P-75%)": r.get("ATH_Profit_Grade", ""),
            "Nifty 500 [OTP] (P-82%)": r.get("Nifty_Alpha_Grade", ""),
            "Sector Index (P-82%)": r.get("Sector_Alpha_Grade", ""),
            "REVISION SCORE": r.get("Review_Score", ""),
            "ACTION": r.get("ACTION", ""),
        }

    enriched_rows = []
    for _, row in closed_journal_df.iterrows():
        inst = str(row.get("symbol", row.get("Instrument", ""))).replace(".NS", "").upper()
        d = row.to_dict()
        scores = scored_map.get(
            inst,
            {
                "Market CAP": "",
                "ATH (P-66%)": "",
                "ATP (P-75%)": "",
                "Nifty 500 [OTP] (P-82%)": "",
                "Sector Index (P-82%)": "",
                "REVISION SCORE": "",
                "ACTION": "",
            },
        )
        d.update(scores)
        enriched_rows.append(d)

    return pd.DataFrame(enriched_rows)


def build_universe_750_scores_df(
    df_scored_results: pd.DataFrame,
    temp_avg_cost_map: dict,
    temp_entry_date: str = "2026-10-25",
    total_capital: float = 150000.0,
    allocation_pct: float = 0.065,
    open_journal_df: pd.DataFrame = None,
) -> pd.DataFrame:
    """Transforms df_scored_results into Universe-of-750-Stocks-Scores layout."""
    if df_scored_results is None or df_scored_results.empty:
        return pd.DataFrame()

    df = df_scored_results.copy()
    base_allowed_inv_per_stock = allocation_pct * total_capital
    entry_dt = pd.to_datetime(temp_entry_date)
    today_dt = pd.Timestamp.now()

    open_symbols = set()
    if open_journal_df is not None and not open_journal_df.empty:
        for _, r in open_journal_df.iterrows():
            open_symbols.add(str(r["Instrument"]).replace(".NS", "").upper())

    open_metrics_map = {}
    if open_journal_df is not None and not open_journal_df.empty:
        for _, row in open_journal_df.iterrows():
            inst = str(row["Instrument"]).replace(".NS", "").upper()
            open_metrics_map[inst] = {
                "Allowed Investment": row.get(
                    "Allowed Investment", base_allowed_inv_per_stock
                ),
                "Allowed Shares": row.get("Allowed Shares", 0),
                "Qty.": row.get("Qty.", ""),
                "Avg. cost": row.get("Avg. cost", ""),
                "Invested": row.get("Invested", ""),
                "Cur. val": row.get("Cur. val", ""),
                "P&L": row.get("P&L", ""),
                "Net chg.": row.get("Net chg.", ""),
                "Day chg.": row.get("Day chg.", ""),
                "Defined SL [RISK]": row.get("Defined SL [RISK]", ""),
                "Invested Allocation %": row.get("Invested Allocation %", ""),
                "Current Value Allocation %": row.get(
                    "Current Value Allocation %", ""
                ),
                "entry_time": row.get("entry_time", temp_entry_date),
                "holding_duration": row.get("holding_duration", ""),
            }

    universe_rows = []
    for _, row in df.iterrows():
        ticker = row["Ticker"]
        ticker_clean = str(ticker).replace(".NS", "").upper()
        ltp = float(row["Stock_Price"]) if pd.notna(row["Stock_Price"]) else 0.0

        remarks_val = (
            "Already In Trade" if ticker_clean in open_symbols else ""
        )

        if ticker_clean in open_metrics_map:
            m = open_metrics_map[ticker_clean]
            allowed_inv = m["Allowed Investment"]
            allowed_shares = m["Allowed Shares"]
            temp_data = {
                "Allowed Investment": allowed_inv,
                "Allowed Shares": allowed_shares,
                "Temp Entry Date": m["entry_time"],
                "Temp Holding Duration": m["holding_duration"],
                "Temp Qty.": m["Qty."],
                "Temp Avg. cost": m["Avg. cost"],
                "Temp Invested": m["Invested"],
                "Temp Cur. val": m["Cur. val"],
                "Temp P&L": m["P&L"],
                "Temp Net chg.": m["Net chg."],
                "Temp Day chg.": m["Day chg."],
                "Temp Allowed Investment": allowed_inv,
                "Temp Allowed Shares": allowed_shares,
                "Temp Invested Allocation %": m["Invested Allocation %"],
                "Temp Current Value Allocation %": m[
                    "Current Value Allocation %"
                ],
                "Temp Defined SL [RISK]": m["Defined SL [RISK]"],
            }
        else:
            avg_cost = temp_avg_cost_map.get(ticker_clean) or temp_avg_cost_map.get(
                ticker
            )
            if avg_cost is not None and avg_cost > 0 and ltp > 0:
                qty = max(1, int(base_allowed_inv_per_stock / avg_cost))
                invested = round(qty * avg_cost, 3)
                cur_val = round(qty * ltp, 3)
                pnl = round(cur_val - invested, 3)
                net_chg_pct = round(((ltp - avg_cost) / avg_cost) * 100, 3)
                allowed_inv = round(base_allowed_inv_per_stock - invested, 3)
                allowed_shares = (
                    int(allowed_inv / ltp)
                    if allowed_inv > 0 and ltp > 0
                    else 0
                )
                invested_alloc_pct = round(
                    (invested / total_capital) * 100, 3
                )
                cur_val_alloc_pct = round((cur_val / total_capital) * 100, 3)
                risk_amount = (
                    (total_capital * 0.01) / 2
                    if invested < 6000
                    else (total_capital * 0.01)
                )
                defined_sl = round(avg_cost - (risk_amount / qty), 3)
                duration_str = format_holding_duration(today_dt - entry_dt)

                temp_data = {
                    "Allowed Investment": allowed_inv,
                    "Allowed Shares": allowed_shares,
                    "Temp Entry Date": temp_entry_date,
                    "Temp Holding Duration": duration_str,
                    "Temp Qty.": qty,
                    "Temp Avg. cost": round(avg_cost, 3),
                    "Temp Invested": invested,
                    "Temp Cur. val": cur_val,
                    "Temp P&L": pnl,
                    "Temp Net chg.": net_chg_pct,
                    "Temp Day chg.": 0.0,
                    "Temp Allowed Investment": allowed_inv,
                    "Temp Allowed Shares": allowed_shares,
                    "Temp Invested Allocation %": invested_alloc_pct,
                    "Temp Current Value Allocation %": cur_val_alloc_pct,
                    "Temp Defined SL [RISK]": defined_sl,
                }
            else:
                allowed_inv = round(base_allowed_inv_per_stock, 3)
                allowed_shares = int(allowed_inv / ltp) if ltp > 0 else 0
                temp_data = {
                    "Allowed Investment": allowed_inv,
                    "Allowed Shares": allowed_shares,
                    "Temp Entry Date": "",
                    "Temp Holding Duration": "",
                    "Temp Qty.": "",
                    "Temp Avg. cost": "",
                    "Temp Invested": "",
                    "Temp Cur. val": "",
                    "Temp P&L": "",
                    "Temp Net chg.": "",
                    "Temp Day chg.": "",
                    "Temp Allowed Investment": allowed_inv,
                    "Temp Allowed Shares": allowed_shares,
                    "Temp Invested Allocation %": "",
                    "Temp Current Value Allocation %": "",
                    "Temp Defined SL [RISK]": "",
                }

        universe_rows.append(
            {
                "Instrument": row["Ticker"],
                "LTP": row["Stock_Price"],
                "Sector": row["Sector"],
                "Broad Based Index": row["Board_Based"],
                "Market CAP": row["Market_Cap_Tier"],
                "Upcoming Event Date": row["Upcoming_Event_Date"],
                "Upcoming Event Name": row["Upcoming_Event_Name"],
                "Volatility_%": row["Volatility_%"],
                "Sharpe_Ratio": row["Sharpe_Ratio"],
                "ATH (P-66%)": row["ATH_Price_Grade"],
                "ATP (P-75%)": row["ATH_Profit_Grade"],
                "Nifty 500 [OTP] (P-82%)": row["Nifty_Alpha_Grade"],
                "Sector Index (P-82%)": row["Sector_Alpha_Grade"],
                "REVISION SCORE": row["Review_Score"],
                "ACTION": row["ACTION"],
                "Allowed Investment": temp_data["Allowed Investment"],
                "Allowed Shares": temp_data["Allowed Shares"],
                "REMARKS": remarks_val,
                "Temp Entry Date": temp_data["Temp Entry Date"],
                "Temp Holding Duration": temp_data["Temp Holding Duration"],
                "Temp Qty.": temp_data["Temp Qty."],
                "Temp Avg. cost": temp_data["Temp Avg. cost"],
                "Temp Invested": temp_data["Temp Invested"],
                "Temp Cur. val": temp_data["Temp Cur. val"],
                "Temp P&L": temp_data["Temp P&L"],
                "Temp Net chg.": temp_data["Temp Net chg."],
                "Temp Day chg.": temp_data["Temp Day chg."],
                "Temp Allowed Investment": temp_data["Temp Allowed Investment"],
                "Temp Allowed Shares": temp_data["Temp Allowed Shares"],
                "Temp Invested Allocation %": temp_data[
                    "Temp Invested Allocation %"
                ],
                "Temp Current Value Allocation %": temp_data[
                    "Temp Current Value Allocation %"
                ],
                "Temp Defined SL [RISK]": temp_data["Temp Defined SL [RISK]"],
            }
        )

    return pd.DataFrame(universe_rows)


def export_to_excel_direct(
    output_path: str,
    df_scored_results: pd.DataFrame,
    universe_750_df: pd.DataFrame = None,
    open_journal_df: pd.DataFrame = None,
    closed_journal_df: pd.DataFrame = None,
):
    """Creates a stylized multi-tab Excel file directly with custom ACTION colors."""
    logger.info(f"Generating styled Excel report: {output_path} ...")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    header_blue = PatternFill(
        start_color="366092", end_color="366092", fill_type="solid"
    )
    header_navy = PatternFill(
        start_color="1F497D", end_color="1F497D", fill_type="solid"
    )
    header_grey = PatternFill(
        start_color="595959", end_color="595959", fill_type="solid"
    )
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    zebra_fill = PatternFill(
        start_color="F9FAFB", end_color="F9FAFB", fill_type="solid"
    )

    action_styles = {
        "+ADD": {
            "fill": PatternFill(
                start_color="00B050", end_color="00B050", fill_type="solid"
            ),
            "font": Font(name="Calibri", size=10, bold=True, color="FFFFFF"),
        },
        "HOLD": {
            "fill": PatternFill(
                start_color="0070C0", end_color="0070C0", fill_type="solid"
            ),
            "font": Font(name="Calibri", size=10, bold=True, color="FFFFFF"),
        },
        "REVIEW": {
            "fill": PatternFill(
                start_color="FFD966", end_color="FFD966", fill_type="solid"
            ),
            "font": Font(name="Calibri", size=10, bold=True, color="000000"),
        },
        "REPLACE": {
            "fill": PatternFill(
                start_color="ED7D31", end_color="ED7D31", fill_type="solid"
            ),
            "font": Font(name="Calibri", size=10, bold=True, color="FFFFFF"),
        },
        "EXIT": {
            "fill": PatternFill(
                start_color="FF0000", end_color="FF0000", fill_type="solid"
            ),
            "font": Font(name="Calibri", size=10, bold=True, color="FFFFFF"),
        },
    }

    sheets_to_write = [
        ("Open-Partial-Journal", open_journal_df),
        ("Closed-Journal", closed_journal_df),
        ("Universe-of-750-Stocks-Scores", universe_750_df),
        ("Python-750-Computed-Scores", df_scored_results),
    ]

    for sheet_name, df in sheets_to_write:
        if df is None or df.empty:
            continue

        ws = wb.create_sheet(title=sheet_name)
        ws.views.sheetView[0].showGridLines = True

        headers = list(df.columns)
        ws.append(headers)

        for col_idx, col_name in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )

            if "Temp" in col_name:
                cell.fill = header_navy
            elif col_name in ["ACTION", "REMARKS", "REMAKS"]:
                cell.fill = header_grey
            else:
                cell.fill = header_blue

        for row_idx, row in enumerate(df.itertuples(index=False), 2):
            ws.append(list(row))
            is_even = row_idx % 2 == 0

            for col_idx in range(1, len(headers) + 1):
                col_name = headers[col_idx - 1]
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.border = thin_border
                if is_even:
                    cell.fill = zebra_fill

                if col_name == "ACTION" and str(cell.value) in action_styles:
                    cell.fill = action_styles[str(cell.value)]["fill"]
                    cell.font = action_styles[str(cell.value)]["font"]
                    cell.alignment = Alignment(
                        horizontal="center", vertical="center"
                    )

                if isinstance(cell.value, (int, float)):
                    if any(
                        k in col_name
                        for k in [
                            "Price",
                            "cost",
                            "LTP",
                            "P&L",
                            "val",
                            "Invested",
                            "SL",
                            "Investment",
                        ]
                    ):
                        cell.number_format = "#,##0.00"
                    elif "%" in col_name or "Ratio" in col_name or "SCORE" in col_name:
                        cell.number_format = "0.00"

        ws.freeze_panes = "A2"

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 38)

        logger.info(f" -> Created sheet: '{sheet_name}' ({len(df)} rows)")

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    logger.info(f"[Success] Direct Excel report generated and saved to: {output_path}")